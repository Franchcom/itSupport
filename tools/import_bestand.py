#!/usr/bin/env python3
"""Bestandsdaten (Excel/Word) in eine Import-Datei fuer itSupport umwandeln.

    pip install openpyxl python-docx
    python3 tools/import_bestand.py --config tools/zuordnung.json --quelle ~/Kundendaten --ziel itsupport-import.json

Die Zuordnung (welche Datei/welches Blatt/welche Zeilen zu welchem Kunden
gehoeren) steht in einer JSON-Datei, siehe tools/zuordnung.beispiel.json. Die
echte Zuordnung enthaelt Kundennamen und bleibt deshalb lokal (.gitignore).

Die erzeugte Datei enthaelt alle Passwoerter im KLARTEXT. In der App unter
Einstellungen -> Import einspielen und danach sofort loeschen.

Das Skript gibt nie Passwoerter aus, nur Zaehler.
"""

import argparse
import base64
import datetime as dt
import glob
import io
import json
import os
import posixpath
import re
import sys
import zipfile

try:
    import openpyxl
    from openpyxl.utils import column_index_from_string
except ImportError:
    sys.exit('Bitte zuerst installieren: pip install openpyxl python-docx')

CATEGORIES = [
    'Fernwartung', 'Netzwerk', 'Server & NAS', 'Benutzer & Geräte', 'Microsoft 365',
    'E-Mail', 'Portale & Web', 'Backup', 'Notizen',
]

# ---------------------------------------------------------------- Heuristiken

SECRET_LABEL = re.compile(r'(^|[^a-z])(pw|pwd|pa+s+w\w*|kennw\w*|pass|pin|sicherheitsfrage\w*|startpw|dsrm)([^a-z]|$)', re.I)
SECRET_INLINE = re.compile(r'(^|[^a-z])(pw|passwort|kennwort|pass|password)\s*[:=]', re.I)
IP = re.compile(r'^\d{1,3}(\.\d{1,3}){3}\.?(:\d+)?$')
DATE = re.compile(r'^\d{1,2}\.\d{1,2}\.\d{2,4}$')
HOST = re.compile(r'^[a-z0-9-]+(\.[a-z0-9-]+)+\.[a-z]{2,}$', re.I)
EMAIL = re.compile(r'^[^\s@]+@[^\s@]+\.[a-z]{2,}$', re.I)
NOT_SECRET_WORDS = re.compile(r'^(ms)?(office|o|m|microsoft)365$|^win(dows)?\d+$|^server\d+$|^eth\d+$|^plan\d$|^ATU\d{8}$', re.I)
DOMAIN_START = re.compile(r'^([a-z0-9-]+\.)+[a-z]{2,}(?=\s|$)', re.I)


def looks_like_password(token):
    t = token.strip('()[]{},;"\'')
    if len(t) < 6 or len(t) > 48:
        return False
    if IP.match(t) or DATE.match(t) or HOST.match(t) or EMAIL.match(t) or NOT_SECRET_WORDS.match(t):
        return False
    if re.match(r'^https?[:.]//', t, re.I) or t.isdigit():
        return False
    has_alpha = re.search(r'[a-zA-Z]', t)
    return bool(has_alpha and (re.search(r'\d', t) or re.search(r'[?!#+*$%&]', t)))


def is_secret(label, value):
    if not value:
        return False
    if SECRET_LABEL.search(label or '') or SECRET_INLINE.search(value):
        return True
    return any(looks_like_password(tok) for tok in re.split(r'[\s/|=:]+', value))


def text(v):
    """Zellwert als lesbarer Text."""
    if v is None:
        return ''
    if isinstance(v, dt.datetime):
        return v.strftime('%d.%m.%Y')
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace('\xa0', ' ')
    s = '\n'.join(re.sub(r'[ \t]+', ' ', line).strip() for line in s.splitlines())
    return s.strip()


def entry(label, value, secret=None, kind=None, note=''):
    label, value = (label or '').strip().rstrip(':').strip(), (value or '').strip()
    # Steht ein Passwort schon in der Bezeichnung ("user= Geheim1?"), wandert
    # alles in den (verdeckten) Wert, damit nichts im Klartext sichtbar bleibt.
    if label and any(looks_like_password(t) for t in re.split(r'[\s/|=:]+', label)):
        label, value, secret = '', ' · '.join(x for x in (label, value) if x), True
    e = {'label': label, 'value': value, 'secret': is_secret(label, value) if secret is None else bool(secret and value)}
    if kind:
        e['kind'] = kind
    if note:
        e['note'] = note
    return e


# ------------------------------------------------------------ Kundenregister

class Registry:
    def __init__(self):
        self.customers = {}

    def get(self, name):
        if name not in self.customers:
            self.customers[name] = {'name': name, 'note': '', 'sections': []}
        return self.customers[name]

    def note(self, name, line):
        c = self.get(name)
        if line not in c['note']:
            c['note'] = (c['note'] + '\n' + line).strip()

    def add(self, customer, title, category, entries):
        entries = [e for e in entries if e['label'] or e['value']]
        entries = [e for e in entries if e['value'] or e.get('note') or len(entries) == 1]
        if not entries:
            return
        assert category in CATEGORIES, category
        self.get(customer)['sections'].append({'title': title, 'category': category, 'entries': entries})


# ------------------------------------------------------------- Excel-Helfer

def col(letter):
    return column_index_from_string(letter) - 1


def sheet_rows(ws, first, last):
    """{zeilennummer: [zellen als text]} fuer nicht-leere Zeilen im Bereich."""
    out = {}
    for r in range(first, min(last, ws.max_row) + 1):
        cells = [text(c.value) for c in ws[r]]
        if any(cells):
            out[r] = cells
    return out


def free_entries(lines):
    """Freie Zeilen in Eintraege verwandeln.

    lines: Liste von Zellenlisten (nur nicht-leere Zellen).
    - mehrere Zellen:           erste = Bezeichnung, Rest = Wert
    - "Bezeichnung: Wert":      aufgeteilt
    - "Bezeichnung:" allein:    nimmt die folgenden Einzelzeilen ohne ':' als Wert
    - sonst:                    Zeile als Notiz (ohne Bezeichnung)
    """
    out = []
    i = 0
    while i < len(lines):
        cells = lines[i]
        i += 1
        if len(cells) >= 2:
            out.append(entry(cells[0], ' · '.join(cells[1:])))
            continue
        line = cells[0]
        if re.match(r'^https?[:.]//', line, re.I) or ':' not in line:
            out.append(entry('', line))
            continue
        label, _, value = line.partition(':')
        value = value.strip()
        if not value:
            parts = []
            while i < len(lines) and len(lines[i]) == 1 and len(parts) < 6:
                nxt = lines[i][0]
                if nxt.rstrip().endswith(':') or (':' in nxt and not re.match(r'^https?[:.]//', nxt, re.I)):
                    break
                parts.append(nxt)
                i += 1
            value = ' · '.join(parts)
            if not value:
                continue
        out.append(entry(label, value))
    return out


def region(reg, ws, cfg, customer):
    first, last = cfg['rows']
    rows = sheet_rows(ws, first, last)
    mode = cfg.get('mode', 'free')
    category = cfg.get('category', 'Notizen')
    title = cfg.get('title', '')

    if mode == 'free':
        reg.add(customer, title, category, free_entries([[c for c in cells if c] for cells in rows.values()]))
        return

    header_row = cfg['header_row']
    headers = [text(c.value) for c in ws[header_row]]
    for letter, label in cfg.get('header_override', {}).items():
        while len(headers) <= col(letter):
            headers.append('')
        headers[col(letter)] = label
    data = {r: cells for r, cells in rows.items() if r != header_row}

    if mode == 'kv':
        reg.add(customer, title, category, [entry(c[0], ' · '.join(c[1:])) for c in ([x for x in cells if x] for cells in data.values())])
        return

    if mode == 'rows':
        out = []
        for cells in data.values():
            filled = [(i, v) for i, v in enumerate(cells) if v]
            name = filled[0][1]
            rest = ' · '.join(f'{headers[i]}: {v}' if i < len(headers) and headers[i] else v for i, v in filled[1:])
            out.append(entry(name, rest))
        reg.add(customer, title, category, out)
        return

    if mode == 'cards':
        title_cols = [col(x) for x in cfg.get('title_cols', ['A'])]
        min_cells = cfg.get('min_cells', 3)
        leftovers = []
        for cells in data.values():
            filled = [(i, v) for i, v in enumerate(cells) if v]
            if len(filled) < min_cells:
                leftovers.append([v for _, v in filled])
                continue
            tcol = next((i for i in title_cols if i < len(cells) and cells[i]), filled[0][0])
            card_title = cells[tcol]
            out = []
            for i, v in filled:
                if i == tcol:
                    continue
                label = headers[i] if i < len(headers) else ''
                kind = None
                if 'teamviewer' in label.lower() and v.replace(' ', '').isdigit():
                    kind = 'teamviewer'
                if label.strip().lower() == 'admin':
                    out.append(entry(label, v, secret=True))
                else:
                    out.append(entry(label, v, kind=kind))
            cat = category
            alt = cfg.get('category_if_empty')
            if alt and not cells[col(alt['col'])]:
                cat = alt['category']
            reg.add(customer, card_title, cat, out)
        if leftovers:
            reg.add(customer, f'{title} – Weitere' if title else 'Weitere', category, free_entries(leftovers))
        return

    raise SystemExit(f'Unbekannter Modus: {mode}')


# ------------------------------------------------------------- AnyDesk-Liste

def import_anydesk(reg, ws, cfg):
    c = {k: (col(v) if isinstance(v, str) else [col(x) for x in v]) for k, v in cfg['cols'].items()}
    seen = set()
    count = 0
    rows = sheet_rows(ws, cfg['first_row'], ws.max_row)
    named_ids = {cells[c['anydesk']].replace(' ', '') for cells in rows.values()
                 if len(cells) > c['anydesk'] and cells[c['name']] and cells[c['anydesk']]}
    for r, cells in rows.items():
        get = lambda i: cells[i] if i < len(cells) else ''
        name = get(c['name'])
        ad = get(c['anydesk']).replace(' ', '')
        tv = get(c['teamviewer']).replace(' ', '')
        pws = [get(i) for i in c['password'] if get(i)]
        if not ad and not tv:
            continue
        if ad in seen or (not name and ad in named_ids):
            continue
        seen.add(ad)
        customer, device = cfg.get('fallback', 'Nicht zugeordnet'), name
        for prefix, target in cfg['map']:
            if name.lower().startswith(prefix.lower()):
                customer = target
                break
        if not name:
            # Unbeschriftete Zeile: ein nicht-geheimer Text in der Passwortspalte ist eher ein Geraetename.
            label_like = [p for p in pws if not looks_like_password(p)]
            device = label_like[0] if label_like else 'Unbenanntes Gerät'
            pws = [p for p in pws if p not in label_like]
        entries = []
        if ad:
            entries.append(entry('AnyDesk', ad, secret=False, kind='anydesk'))
        if tv:
            entries.append(entry('TeamViewer', tv, secret=False, kind='teamviewer'))
        for p in pws:
            entries.append(entry('Anmelde-Passwort', p, secret=True))
        reg.add(customer, device, 'Fernwartung', entries)
        count += 1
    for customer, line in cfg.get('notes', {}).items():
        reg.note(customer, line)
    return count


# --------------------------------------------------------- Microsoft 365

def import_o365(reg, ws, cfg):
    c = {k: col(v) for k, v in cfg['cols'].items()}
    domain_map = {k.lower(): v for k, v in cfg.get('domain_map', {}).items()}
    only = {d.lower() for d in cfg['only_domains']} if cfg.get('only_domains') else None
    stop = [re.compile(p) for p in cfg.get('stop_patterns', [])]
    aliases = [(re.compile(p), d.lower()) for p, d in cfg.get('block_aliases', [])]
    cost_cols = [col(x) for x in cfg.get('cost_cols', [])]
    cost_header = [text(ws.cell(row=cfg['cost_header_row'], column=i + 1).value) for i in cost_cols] if cost_cols else []
    blocks = []
    block = None

    for r, cells in sheet_rows(ws, *cfg['rows']).items():
        get = lambda k: cells[c[k]] if k in c and c[k] < len(cells) else ''
        a = get('admin')
        if a and any(p.search(a) for p in stop):
            block = None
            continue
        alias = next((d for p, d in aliases if a and p.search(a)), None)
        m = DOMAIN_START.match(a) if a and '@' not in a.split()[0] and not alias else None
        if m or alias:
            domain = alias or m.group(0).lower()
            block = {'domain': domain, 'entries': [], 'logins': set(), 'last': None}
            blocks.append(block)
            rest = (a if alias else a[m.end():]).strip(' -–')
            if rest:
                block['entries'].append(entry('Hinweis', rest))
            a = ''
        if block is None:
            continue
        E = block['entries']
        if a:
            if EMAIL.match(a):
                E.append(entry('Tenant-Admin', a, secret=False, kind='email'))
            elif re.match(r'^https?://', a, re.I):
                E.append(entry('Admin-Portal', a, secret=False, kind='url'))
            elif re.match(r'^(user|pw|passwort)\s*:', a, re.I):
                label, _, value = a.partition(':')
                E.append(entry('Admin-Passwort' if label.strip().lower() in ('pw', 'passwort') else 'Tenant-Admin', value))
            elif looks_like_password(a.split()[0]):
                E.append(entry('Admin-Passwort', a, secret=True))
            else:
                E.append(entry('Hinweis', a))
        mailbox = get('mailbox')
        if mailbox and '@' in mailbox:
            addr, _, extra = mailbox.partition('/')
            pw = get('password')
            value = ' / '.join(x.strip() for x in (pw, extra) if x.strip())
            info = {
                'Lizenz': get('license'),
                'Benutzername': get('username') if get('username') and get('username') != addr.strip() else '',
                'Backup': ' '.join(x for x in (get('backup_date'), get('backup_place')) if x),
                'Status': get('status'),
            }
            e = entry(addr.strip(), value, secret=True, kind=None)
            e['_aliases'] = [x for x in re.split(r'\s+', get('aliases')) if x]
            e['_info'] = info
            E.append(e)
            block['last'] = e
        elif mailbox:
            E.append(entry('Info', mailbox))
        elif get('aliases') and block['last'] is not None:
            block['last']['_aliases'] += [x for x in re.split(r'\s+', get('aliases')) if x]
        elif get('status') and not a:
            E.append(entry('Status', get('status'), secret=False))
        costs = [(h, cells[i]) for h, i in zip(cost_header, cost_cols) if i < len(cells) and cells[i]]
        if len(costs) >= 3:
            E.append(entry('Lizenzkosten', ' · '.join(f'{h}: {v.replace(".", ",")}' for h, v in costs), secret=False))
        tu, tp = get('tenant_user'), get('tenant_pw')
        if tu and (tu, tp) not in block['logins']:
            block['logins'].add((tu, tp))
            E.append(entry(f'O365-Anmeldung {tu}', tp, secret=True))

    count = 0
    for b in blocks:
        if only is not None and b['domain'] not in only:
            continue
        customer = domain_map.get(b['domain'], b['domain'])
        for e in b['entries']:
            aliases = e.pop('_aliases', None)
            info = e.pop('_info', None)
            if info is not None:
                parts = [f'{k}: {v}' for k, v in info.items() if v]
                if aliases:
                    parts.append('Aliase: ' + ', '.join(dict.fromkeys(aliases)))
                e['note'] = ' · '.join(parts)
        suffix = f' ({cfg["title_suffix"]})' if cfg.get('title_suffix') else ''
        reg.add(customer, f'Microsoft 365 – {b["domain"]}{suffix}', 'Microsoft 365', b['entries'])
        count += 1
    return count


# ------------------------------------------------------------------ Word

def import_docx(reg, path, cfg):
    try:
        import docx
    except ImportError:
        sys.exit('Bitte zuerst installieren: pip install python-docx')
    d = docx.Document(path)
    lines = []
    for p in d.paragraphs:
        for line in p.text.splitlines():
            line = re.sub(r'[ \t\xa0]+', ' ', line).strip()
            if line:
                lines.append(line)
    drop = [re.compile(x, re.I) for x in cfg.get('drop', [])]
    lines = [l for l in lines if not any(p.search(l) for p in drop)]

    starts = cfg.get('sections', [])
    current = {'title': cfg.get('title', 'Dokumentation'), 'category': cfg.get('category', 'Notizen'), 'lines': []}
    groups = [current]
    for line in lines:
        hit = next((s for s in starts if line.lower().startswith(s['start'].lower())), None)
        if hit:
            current = {'title': hit.get('title', hit['start']), 'category': hit.get('category', 'Notizen'), 'lines': []}
            groups.append(current)
            rest = line[len(hit['start']):].strip()
            if hit.get('keep_line'):
                current['lines'].append([line])
            elif rest.lstrip(':').strip():
                current['lines'].append(['Info' + (rest if rest.startswith(':') else ': ' + rest)])
            continue
        current['lines'].append([line])
    for g in groups:
        reg.add(cfg['customer'], g['title'], g['category'], free_entries(g['lines']))
    images = len(d.inline_shapes)
    if images:
        reg.note(cfg['customer'], f'Word-Dokument „{os.path.basename(path)}“ enthält {images} Screenshots, die nicht importiert wurden.')


# ------------------------------------------------------------------ Bilder

def xlsx_sheet_images(path, sheet_name):
    """Eingebettete Bilder eines Tabellenblatts: [(name, zeile, bytes)], nach Zeile sortiert."""
    z = zipfile.ZipFile(path)
    read = lambda n: z.read(n).decode('utf-8')
    rels = lambda n: dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', read(n))) if n in z.namelist() else {}
    wb = read('xl/workbook.xml')
    rid = next((r for n, r in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wb) if n == sheet_name), None)
    if not rid:
        raise SystemExit(f'Blatt „{sheet_name}“ nicht gefunden in {path}')
    sheet = posixpath.normpath(posixpath.join('xl', rels('xl/_rels/workbook.xml.rels')[rid].lstrip('/').replace('xl/', '', 1)))
    out = []
    srels = rels(posixpath.join(posixpath.dirname(sheet), '_rels', posixpath.basename(sheet) + '.rels'))
    for target in srels.values():
        if 'drawing' not in target:
            continue
        drawing = posixpath.normpath(posixpath.join(posixpath.dirname(sheet), target))
        drels = rels(posixpath.join(posixpath.dirname(drawing), '_rels', posixpath.basename(drawing) + '.rels'))
        for anchor in re.findall(r'<xdr:(?:twoCellAnchor|oneCellAnchor)[\s\S]*?</xdr:(?:twoCellAnchor|oneCellAnchor)>', read(drawing)):
            emb = re.search(r'r:embed="([^"]+)"', anchor)
            row = re.search(r'<xdr:row>(\d+)</xdr:row>', anchor)
            if not emb or emb.group(1) not in drels:
                continue
            media = posixpath.normpath(posixpath.join(posixpath.dirname(drawing), drels[emb.group(1)]))
            out.append((posixpath.splitext(posixpath.basename(media))[0], int(row.group(1)) + 1 if row else 0, z.read(media)))
    return sorted(out, key=lambda x: x[1])


def docx_images(path):
    z = zipfile.ZipFile(path)
    doc = z.read('word/document.xml').decode('utf-8')
    rels = dict(re.findall(r'Id="([^"]+)"[^>]*Target="(media/[^"]+)"', z.read('word/_rels/document.xml.rels').decode('utf-8')))
    out, seen = [], set()
    for rid in re.findall(r'r:embed="([^"]+)"', doc):
        if rid in rels and rels[rid] not in seen:
            seen.add(rels[rid])
            out.append((posixpath.splitext(posixpath.basename(rels[rid]))[0], 0, z.read('word/' + rels[rid])))
    return out


def compress(data, max_side=2000):
    """Auf max. 2000 px verkleinern und als JPEG speichern (wie die App)."""
    try:
        from PIL import Image
    except ImportError:
        sys.exit('Fuer Bilder bitte installieren: pip install pillow')
    im = Image.open(io.BytesIO(data))
    if im.mode in ('RGBA', 'LA', 'P'):
        im = im.convert('RGBA')
        bg = Image.new('RGB', im.size, 'white')
        bg.paste(im, mask=im.split()[-1])
        im = bg
    else:
        im = im.convert('RGB')
    im.thumbnail((max_side, max_side))
    buf = io.BytesIO()
    im.save(buf, 'JPEG', quality=85, optimize=True)
    return buf.getvalue(), im.width, im.height


def import_images(reg, path, cfg):
    found = xlsx_sheet_images(path, cfg['sheet']) if 'sheet' in cfg else docx_images(path)
    titles = cfg.get('titles', {})
    groups = cfg.get('groups', {})
    skip = set(cfg.get('skip', []))
    customer = reg.get(cfg['customer'])
    customer.setdefault('attachments', [])
    n = 0
    for name, row, data in found:
        if name in skip:
            continue
        jpeg, w, h = compress(data)
        where = f"{cfg.get('sheet') or os.path.basename(path)}{f', Zeile {row}' if row else ''}"
        customer['attachments'].append({
            'title': titles.get(name) or f'Bild aus {where}',
            'group': groups.get(name, cfg.get('group', '')),
            'mime': 'image/jpeg',
            'w': w,
            'h': h,
            'data': base64.b64encode(jpeg).decode('ascii'),
        })
        n += 1
    return n


# ------------------------------------------------------------------ Start

def find(base, pattern):
    hits = sorted(glob.glob(os.path.join(base, pattern)))
    if not hits:
        raise SystemExit(f'Keine Datei gefunden für „{pattern}“ in {base}')
    return hits[0]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--config', required=True)
    ap.add_argument('--quelle', required=True, help='Ordner mit den Excel/Word-Dateien')
    ap.add_argument('--ziel', default='itsupport-import.json')
    args = ap.parse_args()

    with open(args.config, encoding='utf-8') as f:
        cfg = json.load(f)
    reg = Registry()
    books = {}

    def sheet(pattern, name):
        path = find(args.quelle, pattern)
        if path not in books:
            books[path] = openpyxl.load_workbook(path, data_only=True)
        return books[path][name]

    stats = {}
    for a in cfg.get('anydesk', []):
        stats['AnyDesk-Geräte'] = stats.get('AnyDesk-Geräte', 0) + import_anydesk(reg, sheet(a['file'], a['sheet']), a)
    for s in cfg.get('sheets', []):
        ws = sheet(s['file'], s['sheet'])
        for r in s['regions']:
            region(reg, ws, r, r.get('customer', s['customer']))
    for o in cfg.get('o365', []):
        stats['M365-Domains'] = stats.get('M365-Domains', 0) + import_o365(reg, sheet(o['file'], o['sheet']), o)
    for d in cfg.get('docx', []):
        import_docx(reg, find(args.quelle, d['file']), d)
    for im in cfg.get('images', []):
        stats['Bilder'] = stats.get('Bilder', 0) + import_images(reg, find(args.quelle, im['file']), im)
    for name, line in cfg.get('notes', {}).items():
        reg.note(name, line)

    customers = sorted(reg.customers.values(), key=lambda c: c['name'].lower())
    for c in customers:
        order = {cat: i for i, cat in enumerate(CATEGORIES)}
        c['sections'].sort(key=lambda s: order[s['category']])

    out = {'format': 'itsupport-import', 'version': 1, 'createdAt': dt.datetime.now().isoformat(timespec='seconds'), 'customers': customers}
    with open(args.ziel, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    try:
        os.chmod(args.ziel, 0o600)
    except OSError:
        pass

    for c in customers:
        imgs = len(c.get('attachments', []))
        if imgs:
            c['note'] = re.sub(r'Word-Dokument „[^“]*“ enthält \d+ Screenshots, die nicht importiert wurden\.\n?', '', c['note']).strip()
    n_entries = sum(len(s['entries']) for c in customers for s in c['sections'])
    n_secret = sum(e['secret'] for c in customers for s in c['sections'] for e in s['entries'])
    print(f'{len(customers)} Kunden, {n_entries} Einträge ({n_secret} davon geheim) -> {args.ziel}')
    for k, v in stats.items():
        print(f'  {k}: {v}')
    for c in customers:
        print(f'  {c["name"]}: {len(c["sections"])} Bereiche, {sum(len(s["entries"]) for s in c["sections"])} Einträge')
    print('ACHTUNG: Die Datei enthält Passwörter im Klartext. Nach dem Import löschen.')


if __name__ == '__main__':
    main()
